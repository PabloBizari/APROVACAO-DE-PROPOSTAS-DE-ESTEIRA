import tkinter as tk
from tkinter import filedialog, ttk, messagebox
from tkinter.scrolledtext import ScrolledText
from pathlib import Path
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
from datetime import datetime
import unicodedata
import threading
import queue
import sys
import time

URL = "URL DO SITE"
USUARIO = "SEU_USUARIO"
SENHA = "SUA-SENHA"
ARQUIVO_PADRAO = "propostas.xlsx"

Q = queue.Queue()
ARQUIVO_SELECIONADO = None
PROGRESSO_TOTAL = 0
PROGRESSO_ATUAL = 0


def base_dir():
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def normalizar(txt):
    if txt is None:
        return ""
    txt = str(txt).replace('\xa0', ' ').strip()
    txt = unicodedata.normalize('NFKD', txt)
    txt = ''.join(c for c in txt if not unicodedata.combining(c))
    return ' '.join(txt.split()).lower()


def log(msg):
    Q.put(msg)


def planilha_caminho():
    return Path(ARQUIVO_SELECIONADO) if ARQUIVO_SELECIONADO else base_dir() / ARQUIVO_PADRAO


def abrir_planilha(caminho):
    wb = load_workbook(caminho)
    ws = wb.active
    headers = {normalizar(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}
    
    if 'proposta' not in headers:
        raise Exception("A planilha precisa ter a coluna 'proposta' na linha 1.")
        
    novas_colunas = ['fila_anterior', 'fila_atual', 'status', 'data_hora']
    for col in novas_colunas:
        if col not in headers:
            ws.cell(1, ws.max_column + 1).value = col
            headers[col] = ws.max_column
            
    return wb, ws, headers


def salvar_planilha(wb, caminho):
    for _ in range(10):
        try:
            wb.save(caminho)
            return
        except PermissionError:
            time.sleep(1)
    raise PermissionError(f'Feche o Excel antes de rodar: {caminho}')


def fazer_login(page):
    log('Abrindo sistema...')
    page.goto(URL, wait_until='load')
    page.locator('#EUsuario_CAMPO').fill(USUARIO)
    page.locator('#ESenha_CAMPO').fill(SENHA)
    page.locator('#lnkEntrar').click()
    page.locator('text=NEGOCIAÇÃO').wait_for(timeout=20000)
    log('Login realizado.')


def abrir_menu(page):
    try:
        if page.locator('#ctl00_Cph_AprCons_txtPesquisa_CAMPO').is_visible():
            return
    except:
        pass

    try:
        page.locator('text=NEGOCIAÇÃO').click(timeout=3000)
    except:
        pass
    
    try:
        page.locator("a.divTit", has_text="Autorizador").click(timeout=5000)
        esteira = page.locator('a.dropdown-toggle', has_text='Esteira')
        try:
            esteira.hover(timeout=2000)
        except:
            esteira.click(timeout=2000)
        page.locator('#WFP2010_PWTCNPROP').click(timeout=5000)
        page.locator('#ctl00_Cph_AprCons_txtPesquisa_CAMPO').wait_for(timeout=8000)
    except Exception as e:
        log("Aviso: Menu inacessível. Recarregando página...")
        try:
            page.reload()
            page.wait_for_timeout(2000)
        except:
            pass


def formatar_proposta(proposta):
    s = str(proposta).strip()
    if s.isdigit() and len(s) < 9:
        return s.zfill(9)
    return s


def normalizar_proposta_valor(v):
    s = str(v or '').strip()
    s = ''.join(ch for ch in s if ch.isdigit())
    if s and len(s) < 9:
        s = s.zfill(9)
    return s


def consultar(page, proposta):
    proposta = formatar_proposta(proposta)
    campo = page.locator('#ctl00_Cph_AprCons_txtPesquisa_CAMPO')
    botao = page.locator('#btnPesquisar_txt')
    campo.fill('')
    campo.fill(proposta)
    botao.click()

    try:
        page.wait_for_load_state('networkidle', timeout=1500)
    except:
        pass

    prazo_total_ms = 4000
    passo_ms = 200
    decorrido = 0

    while decorrido < prazo_total_ms:
        try:
            tabela = localizar_grade(page)
            linhas = tabela.locator('tr')
            for i in range(linhas.count()):
                linha = linhas.nth(i)
                txt = normalizar_proposta_valor(linha.text_content())
                if proposta and proposta in txt:
                    return
        except:
            pass
        page.wait_for_timeout(passo_ms)
        decorrido += passo_ms

    page.wait_for_timeout(200)


def localizar_grade(page):
    loc = page.locator('table#ctl00_Cph_AprCons_grdConsulta')
    if loc.count() > 0:
        return loc.first
    tables = page.locator('table')
    for i in range(tables.count()):
        tbl = tables.nth(i)
        ths = tbl.locator('th')
        if ths.count() > 0:
            head = ' '.join(normalizar(ths.nth(j).text_content()) for j in range(ths.count()))
            if 'atividade' in head:
                return tbl
    raise Exception('Não foi possível localizar a grade com a coluna atividade.')


def cabecalhos_grade(tabela):
    ths = tabela.locator('th')
    cols = {}
    for j in range(ths.count()):
        txt = normalizar(ths.nth(j).text_content())
        if txt:
            cols[txt] = j
    return cols


def texto_limpo(s):
    return ' '.join((s or '').replace('\xa0', ' ').split()).strip()


def is_valid_queue(text):
    t = normalizar(text)
    if not t:
        return False
    if t in {'pesquisa por:', 'ordenar por:', 'atividade'}:
        return False
    if t.startswith('pesquisa por:'):
        return False
    if len(t) < 3:
        return False
    return True


def extrair_fila_da_celula_com_link(celula):
    anchors = celula.locator('a')
    candidatos = []
    for i in range(anchors.count()):
        a = anchors.nth(i)
        txt = texto_limpo(a.text_content())
        if not txt:
            continue
        href = a.get_attribute('href') or ''
        onclick = a.get_attribute('onclick') or ''
        aid = a.get_attribute('id') or ''
        score = 0
        if 'Atividade$0' in href:
            score += 4
        if 'getRowJsonFromLink' in onclick:
            score += 4
        if 'ctl02_ctl02' in aid:
            score += 3
        if txt.lower().startswith('pesquisa por:'):
            score -= 100
        if txt.lower().startswith('ordenar por:'):
            score -= 100
        candidatos.append((score, txt, a))
        
    candidatos.sort(key=lambda x: x[0], reverse=True)
    
    for score, txt, a in candidatos:
        if is_valid_queue(txt):
            return txt, a
            
    txt = texto_limpo(celula.text_content())
    if is_valid_queue(txt):
        return txt, None
    return '', None


def achar_linha_por_proposta(tabela, cols, proposta):
    idx_prop = None
    for k in ('nr proposta', 'nrproposta', 'nr. proposta'):
        if k in cols:
            idx_prop = cols[k]
            break
    linhas = tabela.locator('tr')
    alvo = formatar_proposta(proposta)
    for i in range(linhas.count()):
        linha = linhas.nth(i)
        tds = linha.locator('td')
        if tds.count() == 0:
            continue
        valores = [normalizar_proposta_valor(tds.nth(j).text_content()) for j in range(tds.count())]
        if idx_prop is not None and idx_prop < len(valores) and valores[idx_prop] == alvo:
            return linha
        if alvo in valores:
            return linha
    return None


def ler_atividade(page, proposta):
    tabela = localizar_grade(page)
    cols = cabecalhos_grade(tabela)
    if 'atividade' not in cols:
        raise Exception(f"Cabeçalho 'atividade' não encontrado.")
    linha = achar_linha_por_proposta(tabela, cols, proposta)
    if linha is None:
        raise Exception(f'Proposta {proposta} não encontrada na grade.')
    cel = linha.locator('td').nth(cols['atividade'])
    
    atividade, link_elemento = extrair_fila_da_celula_com_link(cel)
    if not atividade:
        raise Exception(f'Não foi possível extrair a atividade da proposta {proposta}.')
        
    return atividade, link_elemento


def clicar_aprovar(page):
    btn_aprova = page.locator('#BBApr_txt')
    btn_aprova.wait_for(state='visible', timeout=8000)
    btn_aprova.click()
    try:
        page.wait_for_load_state('networkidle', timeout=3000)
    except:
        pass


def processar():
    global PROGRESSO_TOTAL, PROGRESSO_ATUAL
    browser = None
    try:
        entrada = planilha_caminho()
        if not entrada.exists():
            raise Exception(f'Arquivo não encontrado: {entrada}')
        wb, ws, headers = abrir_planilha(entrada)
        linhas = [r for r in range(2, ws.max_row + 1) if ws.cell(r, headers['proposta']).value not in (None, '')]
        PROGRESSO_TOTAL = len(linhas)
        PROGRESSO_ATUAL = 0
        log(f'Planilha: {entrada.name} | Propostas: {PROGRESSO_TOTAL}')
        
        if PROGRESSO_TOTAL == 0:
            log('Nenhuma proposta encontrada para processar.')
            root.after(0, lambda: messagebox.showinfo("Aviso", "Nenhuma proposta ativa para processar."))
            return

        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            page = browser.new_page(viewport={'width': 1280, 'height': 780})
            
            page.set_default_timeout(15000)
            page.on("dialog", lambda dialog: dialog.accept())
            
            fazer_login(page)
            abrir_menu(page)
            
            for r in linhas:
                proposta = str(ws.cell(r, headers['proposta']).value).strip()
                PROGRESSO_ATUAL += 1
                log(f'[{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL}] Consultando {proposta}...')
                
                fila_anterior = ''
                fila_atual = ''
                status = ''
                
                try:
                    abrir_menu(page)
                    consultar(page, proposta)
                    
                    atividade_inicial, link_elemento = ler_atividade(page, proposta)
                    atividade_inicial_norm = normalizar(atividade_inicial)
                    
                    fila_anterior = atividade_inicial
                    
                    if atividade_inicial_norm == 'confere lastro':
                        if link_elemento:
                            log(f'[{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL}] "CONFERE LASTRO" detectado. Tramitando...')
                            link_elemento.click()
                            
                            clicar_aprovar(page)
                            abrir_menu(page)
                            
                            log(f'[{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL}] Validando alteração de {proposta}...')
                            consultar(page, proposta)
                            
                            try:
                                nova_atividade, _ = ler_atividade(page, proposta)
                                nova_atividade_norm = normalizar(nova_atividade)
                            except Exception:
                                nova_atividade = 'ENVIA PAGAMENTO'
                                nova_atividade_norm = 'envia pagamento'
                            
                            if nova_atividade_norm == 'envia pagamento' or 'pagamento' in nova_atividade_norm:
                                fila_atual = 'ENVIA PAGAMENTO'
                                status = 'OK'
                                log(f'[{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL}] Confirmado: Movido para "ENVIA PAGAMENTO"')
                            else:
                                fila_atual = nova_atividade
                                status = f'Ficou em {nova_atividade}'
                                log(f'[{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL}] {status}')
                        else:
                            fila_atual = atividade_inicial
                            status = 'ERRO: Link não clicável'
                            log(f'[{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL}] {status}')
                    else:
                        fila_atual = atividade_inicial
                        status = atividade_inicial
                        log(f'[{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL}] Fila: {atividade_inicial} (Ignorada)')
                        
                except Exception as e:
                    if not fila_anterior:
                        fila_anterior = 'N/D'
                    fila_atual = 'N/D'
                    status = f'ERRO: {e}'
                    log(f'[{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL}] {status}')
                    
                ws.cell(r, headers['fila_anterior']).value = fila_anterior
                ws.cell(r, headers['fila_atual']).value = fila_atual
                ws.cell(r, headers['status']).value = status
                ws.cell(r, headers['data_hora']).value = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                salvar_planilha(wb, entrada)
                atualizar_interface()
                
            browser.close()
            browser = None
            
        log('Processo finalizado.')
        root.after(0, lambda: messagebox.showinfo("Sucesso", "Processo finalizado com sucesso!"))
        
    except Exception as e:
        log(f'ERRO GERAL: {e}')
        root.after(0, lambda msg=str(e): messagebox.showerror("Erro", f"Ocorreu um erro geral:\n{msg}"))
    finally:
        if browser:
            try:
                browser.close()
            except:
                pass
        root.after(0, lambda: btn_iniciar.config(state='normal'))


def atualizar_interface():
    if PROGRESSO_TOTAL > 0:
        pct = int((PROGRESSO_ATUAL / PROGRESSO_TOTAL) * 100)
        barra['value'] = pct
        lbl_progresso.config(text=f'{PROGRESSO_ATUAL}/{PROGRESSO_TOTAL} propostas')
    root.update_idletasks()


def iniciar():
    btn_iniciar.config(state='disabled')
    log('Iniciando processo...')
    threading.Thread(target=processar, daemon=True).start()


def escolher_arquivo():
    global ARQUIVO_SELECIONADO
    arq = filedialog.askopenfilename(title='Escolher planilha', filetypes=[('Planilhas Excel', '*.xlsx')])
    if arq:
        ARQUIVO_SELECIONADO = arq
        lbl_arquivo.config(text=Path(arq).name)
        log(f'Arquivo selecionado: {arq}')


def poll_queue():
    try:
        while True:
            msg = Q.get_nowait()
            console.insert(tk.END, msg + '\n')
            console.see(tk.END)
    except queue.Empty:
        pass
    atualizar_interface()
    root.after(150, poll_queue)


# ==========================================
# CONFIGURAÇÃO VISUAL DA INTERFACE GRÁFICA
# ==========================================

# Paleta de Cores
BG_SISTEMA = '#F8FAFC'      # Off-white / Slate 50
BG_CARD = '#FFFFFF'         # Branco limpo para o container
COR_PRIMARIA = '#0284C7'    # Azul Sky 600
COR_SECUNDARIA = '#64748B'  # Cinza Slate 500
COR_TEXTO = '#0F172A'       # Slate 900
COR_CONSOLE_BG = '#0F172A'  # Slate Dark para o terminal
COR_CONSOLE_FG = '#38BDF8'  # Texto ciano Sky para leitura confortável

root = tk.Tk()
root.title('NetCapital - Automação')
root.geometry('580x430')
root.configure(bg=BG_SISTEMA)
root.resizable(False, False)

# Tema e customização da Barra de Progresso
style = ttk.Style()
style.theme_use('clam')
style.configure("TProgressbar", 
                thickness=6, 
                troughcolor='#E2E8F0', 
                background=COR_PRIMARIA, 
                bordercolor='#E2E8F0', 
                lightcolor=COR_PRIMARIA, 
                darkcolor=COR_PRIMARIA)

# Container Principal (Visual clean tipo "Card")
painel = tk.Frame(root, bg=BG_CARD, bd=0, padx=20, pady=20)
painel.place(relx=0.03, rely=0.03, relwidth=0.94, relheight=0.94)

# Título
lbl_titulo = tk.Label(painel, text='Aprovação de propostas', font=('Segoe UI', 14, 'bold'), fg=COR_TEXTO, bg=BG_CARD)
lbl_titulo.pack(anchor='w', pady=(0, 2))

# Subtítulo / Instrução
lbl_instrucao = tk.Label(painel, text='Clique em INICIAR para aprovar as propostas da planilha.', font=('Segoe UI', 10), fg=COR_SECUNDARIA, bg=BG_CARD)
lbl_instrucao.pack(anchor='w', pady=(0, 15))

# Seção de Arquivo (Botão plano e Nome da Planilha)
linha_arquivo = tk.Frame(painel, bg=BG_CARD)
linha_arquivo.pack(fill='x', pady=(0, 15))

btn_arquivo = tk.Button(linha_arquivo, text='Escolher planilha', font=('Segoe UI', 9, 'bold'),
                        bg='#F1F5F9', fg='#334155', activebackground='#E2E8F0', activeforeground='#334155',
                        relief='flat', bd=0, padx=12, pady=5, cursor='hand2', command=escolher_arquivo)
btn_arquivo.pack(side='left')

lbl_arquivo = tk.Label(linha_arquivo, text=ARQUIVO_PADRAO, font=('Segoe UI', 9, 'italic'), fg=COR_SECUNDARIA, bg=BG_CARD)
lbl_arquivo.pack(side='left', padx=15)

# Barra de Progresso Fina e Texto Informativo
barra = ttk.Progressbar(painel, orient='horizontal', mode='determinate', style="TProgressbar")
barra.pack(fill='x', pady=(0, 5))

lbl_progresso = tk.Label(painel, text='0/0 propostas', font=('Segoe UI', 9, 'bold'), fg=COR_PRIMARIA, bg=BG_CARD)
lbl_progresso.pack(anchor='w', pady=(0, 12))

# Console Estilo Terminal Dark (Profissional)
console = ScrolledText(painel, height=8, font=('Consolas', 9), 
                       bg=COR_CONSOLE_BG, fg=COR_CONSOLE_FG, 
                       insertbackground='white', relief='flat', bd=0, padx=8, pady=8)
console.pack(fill='both', expand=True)
console.insert(tk.END, 'Pronto. Selecione a planilha ou use o arquivo padrão e clique em INICIAR.\n')

# Rodapé de Ações
rodape = tk.Frame(painel, bg=BG_CARD)
rodape.pack(fill='x', pady=(15, 0))

btn_iniciar = tk.Button(rodape, text='INICIAR', font=('Segoe UI', 10, 'bold'),
                        bg=COR_PRIMARIA, fg='white', activebackground='#0284C7', activeforeground='white',
                        relief='flat', bd=0, width=12, pady=6, cursor='hand2', command=iniciar)
btn_iniciar.pack(side='left')

assinatura = tk.Label(rodape, text='Projetado por Pablo Di Sisto', font=('Segoe UI', 8), fg='#94A3B8', bg=BG_CARD)
assinatura.pack(side='left', padx=15, fill='y')

btn_fechar = tk.Button(rodape, text='FECHAR', font=('Segoe UI', 10, 'bold'),
                       bg='#F1F5F9', fg='#64748B', activebackground='#E2E8F0', activeforeground='#64748B',
                       relief='flat', bd=0, width=10, pady=6, cursor='hand2', command=root.destroy)
btn_fechar.pack(side='right')

root.after(150, poll_queue)
root.mainloop()